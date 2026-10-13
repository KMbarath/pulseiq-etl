"""Compare workbook sheet headers and row counts with PostgreSQL source tables.

Run with ``python scripts/verify_workbook_counts.py`` or
``docker compose run --rm warehouse-load python scripts/verify_workbook_counts.py``.
"""
import os
import re
from pathlib import Path
from openpyxl import load_workbook
import psycopg

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "data" / "03_PulseIQ_Analytics_Star_Schema_Dataset.xlsx"
DB = os.environ["DATABASE_URL"]

def table_name(sheet):
    return "source_" + re.sub(r"[^a-z0-9]+", "_", sheet.lower()).strip("_")

def normalized_columns(header):
    columns = []
    seen = set()
    for index, name in enumerate(header):
        column = re.sub(r"[^a-z0-9]+", "_", str(name or f"column_{index + 1}").lower()).strip("_")
        column = column or f"column_{index + 1}"
        if column in seen:
            column = f"{column}_{index + 1}"
        seen.add(column)
        columns.append(column)
    return columns

def main():
    wb = load_workbook(BOOK, read_only=True, data_only=True)
    mismatches = []
    with psycopg.connect(DB) as conn, conn.cursor() as cur:
        for sheet in wb.sheetnames:
            values = wb[sheet].values
            header = next(values, ())
            while header and all(v is None for v in header):
                header = next(values, ())
            if not header:
                mismatches.append((sheet, "header", "non-empty header", "empty sheet"))
                print(f"{sheet}: invalid empty sheet")
                continue
            expected_columns = normalized_columns(header)
            expected = sum(1 for row in values if any(v is not None for v in row))
            name = table_name(sheet)
            cur.execute('SELECT COUNT(*) FROM "' + name + '"')
            actual = cur.fetchone()[0]
            cur.execute(
                "SELECT column_name FROM information_schema.columns "
                "WHERE table_schema = current_schema() AND table_name = %s "
                "ORDER BY ordinal_position",
                (name,),
            )
            actual_columns = [row[0] for row in cur.fetchall()]
            columns_match = expected_columns == actual_columns
            if expected != actual:
                mismatches.append((sheet, "row_count", expected, actual))
            if not columns_match:
                mismatches.append((sheet, "columns", expected_columns, actual_columns))
            print(
                f"{sheet}: rows workbook={expected}, postgres={actual}; "
                f"columns_match={columns_match} ({len(expected_columns)} columns)"
            )
    if mismatches:
        raise SystemExit(f"Workbook/PostgreSQL mismatches: {mismatches}")
    print("All workbook sheet headers and row counts match PostgreSQL source tables.")

if __name__ == "__main__":
    main()
